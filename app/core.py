"""Persistent exam cards, template-preserving XLSX exchange and review checks."""
from pathlib import Path
import ast
import copy
import hashlib
import io
import json
import math
import os
import re
import threading
import uuid
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime
from functools import lru_cache
import openpyxl
from . import runtime
from .catalog import subject_id, infer_subject, subjects

ROOT = Path(__file__).resolve().parents[1]
DATA = runtime.data_dir(ROOT)
TEMPLATE = ROOT / 'templates' / '문항카드-템플릿.xlsx'
TYPES = ['단일 정답형', '합답형', '불완전 문장형', '짝짓기형']
IMAGE_MODES = ('strict', 'mixed', 'flexible')
LOCK = threading.RLock()
COLS = {2:'no',3:'standard',4:'standard_name',5:'standard_note',6:'difficulty',7:'type',8:'answer',9:'idea',10:'stem',11:'material',12:'question',13:'statements',14:'choices',15:'points',16:'estimate',17:'discrimination',18:'original_no',19:'concepts',20:'explanation'}
TEXT_FIELDS = ['standard','standard_name','standard_note','difficulty','type','idea','stem','material','question','statements','choices','points','concepts','explanation','review_note','image_mode','image_locked','image_creative','ai_feedback','image_text_physics','image_text_math','image_text_korean']
AUTHORING_MODES = ('ai','manual','legacy')
MATERIALS_REQUEST_MAX = 4000
MATERIAL_KINDS = ('image','text','table')
MATERIAL_REQUEST_IMAGE_MODES = ('strict','mixed','flexible')
MATERIAL_IMAGE_FIELDS = ('image_mode','image_locked','image_creative','image_text_physics','image_text_math','image_text_korean','image_text_none','image_plan_version','image_plan_approval','image_labels','labels_initialized','image_scene','image')
MATERIAL_SERVER_FIELDS = ('image_plan_approval','image')

def now():
    return datetime.now().isoformat(timespec='seconds')

def uid():
    return uuid.uuid4().hex[:16]

def normalize_authoring_mode(value):
    if value not in AUTHORING_MODES:
        raise ValueError('작성 방식은 ai, manual, legacy 중 하나여야 합니다.')
    return value

def normalize_materials_request(value):
    if value is None: value=''
    if not isinstance(value,str): raise ValueError('필요한 자료 요청은 문자열이어야 합니다.')
    value=value.strip()
    if len(value)>MATERIALS_REQUEST_MAX: raise ValueError(f'필요한 자료 요청은 {MATERIALS_REQUEST_MAX}자 이하여야 합니다.')
    return value

def normalize_material_request_kinds(value):
    if value is None: value=[]
    if not isinstance(value,list): raise ValueError('필요한 자료 종류는 image, text, table 목록이어야 합니다.')
    normalized=[]
    for kind in value:
        if not isinstance(kind,str): raise ValueError('필요한 자료 종류는 image, text, table 목록이어야 합니다.')
        kind=kind.strip()
        if kind not in MATERIAL_KINDS: raise ValueError('필요한 자료 종류는 image, text, table 중에서 선택해 주세요.')
        if kind not in normalized: normalized.append(kind)
    return normalized

def normalize_material_request_image_mode(value):
    if value is None: value=''
    if not isinstance(value,str): raise ValueError('요청 그림 유형은 문자열이어야 합니다.')
    value=value.strip()
    if value not in ('',*MATERIAL_REQUEST_IMAGE_MODES):
        raise ValueError('요청 그림 유형은 strict / mixed / flexible 중 하나여야 합니다.')
    return value

def validate_material_request_preferences(kinds, image_mode):
    if image_mode and kinds and 'image' not in kinds:
        raise ValueError('그림 유형을 요청하려면 필요한 자료 종류에 image를 포함해 주세요.')
    return kinds,image_mode

def atomic(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix('.tmp')
    temp.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')
    os.replace(temp, path)

def exam_path(eid):
    if not re.fullmatch('[a-f0-9]{16}', eid):
        raise ValueError('회차 ID가 올바르지 않습니다.')
    return DATA / eid

def read_exam(eid):
    p = exam_path(eid) / 'project.json'
    # Keep the file handle lifetime disjoint from our replace-based writer on
    # Windows, where a concurrent reader can make os.replace raise EACCES.
    with LOCK:
        if not p.exists():
            raise ValueError('회차를 찾을 수 없습니다.')
        exam=json.loads(p.read_text(encoding='utf-8'))
    exam.setdefault('subject',infer_subject(exam['cards']))
    for card in exam['cards']:
        card.setdefault('image_text_none',False)
        for key in ('image_text_physics','image_text_math','image_text_korean'): card.setdefault(key,'')
        card.setdefault('calculation_override',None)
        card.setdefault('image_labels',[])
        card.setdefault('image_scene',None)
        card.setdefault('labels_initialized',False)
        card.setdefault('image_plan_version',1)
        for key,value in {'ai_requested':False,'ai_feedback':'','image_mode':'','image_locked':'','image_creative':'','image_plan_approval':None,
                          'authoring_mode':'legacy','materials_request':'','material_request_kinds':[],
                          'material_request_image_mode':''}.items():
            card.setdefault(key,value)
    if not cards_approved(exam): exam['card_approval']=None
    return exam

def save_exam(exam):
    with LOCK:
        for card in exam.get('cards',[]):
            card['authoring_mode']=normalize_authoring_mode(card.get('authoring_mode','legacy'))
            card['materials_request']=normalize_materials_request(card.get('materials_request',''))
            card['material_request_kinds']=normalize_material_request_kinds(card.get('material_request_kinds',[]))
            card['material_request_image_mode']=normalize_material_request_image_mode(card.get('material_request_image_mode',''))
            validate_material_request_preferences(card['material_request_kinds'],card['material_request_image_mode'])
        if (exam.get('card_approval') or {}).get('digest') != cards_digest(exam['cards']):
            exam['card_approval']=None
        exam['updated'] = now()
        exam['revision'] = exam.get('revision', 0) + 1
        atomic(exam_path(exam['id']) / 'project.json', exam)
    return exam

def blank_card(no=1, authoring_mode='legacy'):
    mode=normalize_authoring_mode(authoring_mode)
    return dict(id=uid(), no=no, **{k:'' for k in TEXT_FIELDS}, answer='', estimate='', discrimination='', original_no='', data_kind='none', calculation=False, checks=[], calculation_override=None, review_approved=False, image=None, ai_requested=mode!='manual', authoring_mode=mode, materials_request='', material_request_kinds=[], material_request_image_mode='', image_plan_approval=None,image_labels=[],labels_initialized=False,image_plan_version=2,image_scene=None,image_text_none=False,materials=None)

def legacy_material(card):
    """Project a pre-materials card as its single stable material block."""
    kind=card.get('data_kind','none')
    content=str(card.get('material') or '')
    if kind=='none': return None
    if kind not in MATERIAL_KINDS:
        if content.startswith('(글상자)'): kind='text'
        elif any(line.lstrip().startswith('|') for line in content.splitlines()): kind='table'
        elif content or card.get('image'): kind='image'
        else: kind='none'
    if kind=='none': return None
    block={'id':card['id'],'kind':kind,'label':'','content':content}
    if kind=='image':
        for key in MATERIAL_IMAGE_FIELDS: block[key]=copy.deepcopy(card.get(key))
        block.setdefault('image_mode',''); block.setdefault('image_locked',''); block.setdefault('image_creative','')
        block.setdefault('image_text_physics',''); block.setdefault('image_text_math',''); block.setdefault('image_text_korean','')
        block.setdefault('image_text_none',False); block.setdefault('image_plan_version',2); block.setdefault('image_plan_approval',None)
        block.setdefault('image_labels',[]); block.setdefault('labels_initialized',False); block.setdefault('image_scene',None)
        block.setdefault('image',None)
    return block

def material_blocks(card):
    """Return ordered material blocks, projecting the legacy shape when needed."""
    if card.get('materials') is None:
        block=legacy_material(card)
        return [block] if block else []
    return copy.deepcopy(card.get('materials') or [])

def material_projection(card, material_id=None, require_image=False):
    """Return (projection, material_id) for one block; parent question fields stay shared."""
    blocks=material_blocks(card)
    if material_id is None:
        block=next((b for b in blocks if b.get('kind')=='image'),None) if require_image else (blocks[0] if blocks else None)
    else:
        block=next((b for b in blocks if b.get('id')==material_id),None)
    if block is None or (require_image and block.get('kind')!='image'):
        raise ValueError('대상 자료 블록을 찾을 수 없습니다.' if material_id else '작업할 자료 블록이 없습니다.')
    projected=copy.deepcopy(card)
    projected.update(data_kind=block.get('kind','none'),material=str(block.get('content') or ''))
    if block.get('kind')=='image':
        for key in MATERIAL_IMAGE_FIELDS: projected[key]=copy.deepcopy(block.get(key))
        projected.setdefault('image_mode',''); projected.setdefault('image_locked',''); projected.setdefault('image_creative','')
        projected.setdefault('image_text_physics',''); projected.setdefault('image_text_math',''); projected.setdefault('image_text_korean','')
        projected.setdefault('image_text_none',False); projected.setdefault('image_plan_version',2); projected.setdefault('image_plan_approval',None)
        projected.setdefault('image_labels',[]); projected.setdefault('labels_initialized',False); projected.setdefault('image_scene',None); projected.setdefault('image',None)
    return projected,block['id']

def sync_material_projection(card, projected, material_id):
    """Write one projected block back and refresh compatibility aliases."""
    if card.get('materials') is None:
        card['data_kind']=projected.get('data_kind','none'); card['material']=projected.get('material','')
        if projected.get('data_kind')=='image':
            for key in MATERIAL_IMAGE_FIELDS: card[key]=copy.deepcopy(projected.get(key))
        return card
    blocks=card['materials']
    target=next((b for b in blocks if b.get('id')==material_id),None)
    if target is None: raise ValueError('저장할 자료 블록이 없습니다.')
    target['kind']=projected.get('data_kind',target.get('kind'))
    target['content']=projected.get('material','')
    if target['kind']=='image':
        for key in MATERIAL_IMAGE_FIELDS: target[key]=copy.deepcopy(projected.get(key))
    sync_material_aliases(card)
    return card

def sync_material_aliases(card):
    """Keep old single-material consumers working when the canonical array exists."""
    blocks=card.get('materials') or []
    images=[b for b in blocks if b.get('kind')=='image']
    first=images[0] if images else None
    card['data_kind']='image' if images else (blocks[0].get('kind') if blocks else 'none')
    card['material']=str((blocks[0].get('content') if blocks else '') or '')
    for key in MATERIAL_IMAGE_FIELDS:
        card[key]=copy.deepcopy(first.get(key)) if first else (False if key in ('image_text_none','labels_initialized') else [] if key=='image_labels' else 2 if key=='image_plan_version' else None if key in ('image','image_plan_approval','image_scene') else '')
    return card

def normalize_materials(card):
    """Validate material content and initialize image workflow fields."""
    if card.get('materials') is None:
        card.pop('materials',None)
        return card
    raw=card.get('materials')
    if not isinstance(raw,list) or len(raw)>30: raise ValueError('자료 블록은 30개 이하여야 합니다.')
    from .labels import normalize_labels
    from .vectors import normalize_scene
    blocks=[]; ids=set()
    for source in raw:
        if not isinstance(source,dict): raise ValueError('자료 블록 형식 오류')
        b=copy.deepcopy(source)
        mid=str(b.get('id') or uid())
        if not re.fullmatch('[a-f0-9]{16}',mid) or mid in ids: raise ValueError('자료 블록 ID 오류 또는 중복')
        ids.add(mid); b['id']=mid
        kind=b.get('kind')
        if kind not in MATERIAL_KINDS: raise ValueError('자료 유형은 그림·글상자·표 중 하나여야 합니다.')
        b['label']=str(b.get('label') or '').strip()
        b['content']=str(b.get('content') or '').strip()
        if len(b['label'])>80 or len(b['content'])>50000: raise ValueError('자료 블록 내용이 너무 깁니다.')
        if kind=='image':
            if mid==card.get('id'):
                for key in MATERIAL_IMAGE_FIELDS:
                    if key not in b: b[key]=copy.deepcopy(card.get(key))
            for key,default in (('image_mode',''),('image_locked',''),('image_creative',''),('image_text_physics',''),('image_text_math',''),('image_text_korean','')):
                b[key]=str(b.get(key,default) or '').strip()
            if b['image_mode'] not in ('',*IMAGE_MODES): raise ValueError('그림 유형은 strict / mixed / flexible 중 하나여야 합니다.')
            b['image_text_none']=b.get('image_text_none',False)
            if type(b['image_text_none']) is not bool: raise ValueError('글씨 없음 설정은 Yes 또는 No여야 합니다.')
            b['image_plan_version']=b.get('image_plan_version',2)
            if b['image_plan_version'] not in (1,2): raise ValueError('그림 계획 버전 오류')
            b['image_labels']=normalize_labels(b.get('image_labels',[]))
            b['labels_initialized']=bool(b.get('labels_initialized',False))
            b['image_scene']=normalize_scene(b.get('image_scene'))
            b.setdefault('image_plan_approval',None); b.setdefault('image',None)
        else:
            for key in MATERIAL_IMAGE_FIELDS: b.pop(key,None)
        blocks.append(b)
    card['materials']=blocks
    return sync_material_aliases(card)

def content_hash(card):
    excluded=('image','review_approved','review_note','no','ai_requested','ai_feedback','authoring_mode','materials_request','material_request_kinds','material_request_image_mode','image_plan_approval','image_labels','labels_initialized','image_scene','materials','calculation_override')
    fields = {k:v for k,v in card.items() if k not in excluded}
    blocks=card.get('materials')
    if blocks is not None and not material_array_matches_legacy(card,blocks):
        fields['materials']=[{k:v for k,v in block.items() if k not in ('image','image_plan_approval','image_labels','labels_initialized','image_scene')} for block in blocks]
    return hashlib.sha256(json.dumps(fields, sort_keys=True, ensure_ascii=False).encode()).hexdigest()

def material_array_matches_legacy(card, blocks):
    """A first migration of an unchanged legacy card is a representation-only change."""
    if len(blocks)!=1: return False
    block=blocks[0]
    if block.get('id')!=card.get('id') or block.get('label','') or block.get('kind')!=card.get('data_kind') or block.get('content','')!=card.get('material',''): return False
    if block.get('kind')=='image':
        for key in MATERIAL_IMAGE_FIELDS:
            if key in ('image','image_plan_approval','image_labels','labels_initialized','image_scene'): continue
            if block.get(key)!=card.get(key): return False
    return True

def image_context(card):
    return hashlib.sha256(json.dumps({k:card.get(k) for k in ('idea','stem','material','question','statements','choices','data_kind','image_mode','image_locked','image_creative')},sort_keys=True,ensure_ascii=False).encode()).hexdigest()

def image_plan_errors(card):
    if card['data_kind']!='image': return []
    errors=[]
    if card.get('image_mode') not in IMAGE_MODES: errors.append('그림 유형을 선택해 주세요.')
    mode=card.get('image_mode'); locked=card.get('image_locked','').strip(); creative=card.get('image_creative','').strip()
    if mode in ('strict','mixed') and not locked: errors.append('SVG 규칙을 적어 주세요.')
    if mode in ('flexible','mixed') and not creative: errors.append('창작 허용 영역을 적어 주세요.')
    if mode=='strict' and creative: errors.append('strict의 창작 허용 영역은 빈칸이어야 합니다.')
    if mode=='flexible' and locked: errors.append('flexible의 SVG 규칙은 빈칸이어야 합니다.')
    if mode=='mixed':
        left=element_ids(locked); right=element_ids(creative)
        if not left or not right: errors.append('mixed는 각 줄을 요소ID: 설명 형식으로 작성하세요. 예: floor: 수평 바닥선')
        overlap=left & right
        if overlap: errors.append('두 영역에 같은 요소를 지정할 수 없습니다: '+', '.join(sorted(overlap)))
    text_prompts=[str(card.get(k,'')).strip() for k in ('image_text_physics','image_text_math','image_text_korean')]
    has_text=any(text_prompts)
    if card.get('image_text_none') and (has_text or card.get('image_labels')): errors.append('글씨 없음을 선택했다면 글씨 프롬프트와 객체를 비워 주세요.')
    elif not card.get('image_text_none') and not has_text: errors.append('표시할 글씨가 없으면 “글씨 없음”을 선택해 주세요.')
    if card.get('image_plan_version',1)!=2: errors.append('기존 프롬프트를 그림·글씨 프롬프트로 분리해 주세요.')
    return errors

def element_ids(value):
    lines=[line.strip() for line in value.splitlines() if line.strip()]
    ids=[re.match(r'^([A-Za-z][A-Za-z0-9_-]{0,39})\s*[:：]',line) for line in lines]
    return {m[1].lower() for m in ids} if lines and all(ids) and len({m[1].lower() for m in ids})==len(ids) else set()

def image_plan_context(card):
    return hashlib.sha256(json.dumps([image_context(card),card.get('image_plan_version'),*[card.get(k,'') for k in ('image_text_physics','image_text_math','image_text_korean')],bool(card.get('image_text_none'))],ensure_ascii=False).encode()).hexdigest()

def image_final_digest(card):
    parts=[(card.get('image') or {}).get('sha256'),card.get('image_labels',[])]
    if card.get('image_scene') is not None: parts.append(card['image_scene'])
    return hashlib.sha256(json.dumps(parts,sort_keys=True,ensure_ascii=False).encode()).hexdigest()

def _json_digest(value):
    return hashlib.sha256(json.dumps(value,sort_keys=True,ensure_ascii=False,separators=(',',':')).encode()).hexdigest()

def calculation_input_errors(card):
    """Required expression rows are hard gates; only completed comparisons can be overridden."""
    if not card.get('calculation'):
        return []
    checks=card.get('checks')
    if not isinstance(checks,list) or not checks:
        return ['검산식을 한 개 이상 입력해 주세요.']
    errors=[]
    for index,check in enumerate(checks,1):
        if not isinstance(check,dict):
            errors.append(f'{index}번째 검산식의 형식이 올바르지 않습니다.')
        elif not str(check.get('lhs') or '').strip() or not str(check.get('rhs') or '').strip():
            errors.append(f'{index}번째 검산식의 왼쪽 식과 오른쪽 식을 모두 입력해 주세요.')
    return errors

def calculation_validation_digest(card):
    """Bind an override confirmation to this card and the actual SymPy results."""
    return _json_digest({'content_digest':content_hash(card),'checks':verify_checks(card.get('checks') or [])})

def calculation_override_active(card, record=None):
    record=record if record is not None else card.get('calculation_override')
    if not isinstance(record,dict) or calculation_input_errors(card): return False
    checks=verify_checks(card.get('checks') or [])
    return bool(any(not item['ok'] for item in checks)
                and record.get('content_digest')==content_hash(card)
                and record.get('validation_digest')==calculation_validation_digest(card))

def image_check_results(checks):
    """Image jobs already send server-computed SymPy results; preserve their details."""
    if not isinstance(checks,list) or not checks: return []
    if all(isinstance(item,dict) and type(item.get('ok')) is bool for item in checks):
        return [{'label':str(item.get('label','검산')),'ok':item['ok'],'detail':str(item.get('detail',''))} for item in checks]
    return verify_checks(checks)

def image_validation_digest(card):
    """Digest current content, image source/edit, and server-computed image checks."""
    im=card.get('image') or {}
    return _json_digest({
        'content_digest':content_hash(card),
        'context_digest':image_context(card),
        'plan_digest':image_plan_context(card),
        'image_digest':image_final_digest(card),
        'source_sha256':im.get('sha256'),
        'edited_sha256':im.get('edited_sha256'),
        'edited_digest':im.get('edited_digest'),
        'layer_version':im.get('layer_version'),
        'checks':image_check_results(im.get('checks') or []),
    })

def image_approval_artifact_digest(card):
    im=card.get('image') or {}
    return _json_digest({
        'validation_digest':image_validation_digest(card),
        'final_file':im.get('final_file'),
        'final_sha256':im.get('final_sha256'),
        'final_digest':im.get('final_digest'),
    })

def image_override_active(card, record=None):
    im=card.get('image') or {}
    record=record if record is not None else im.get('approval_override')
    return bool(im.get('approved') and isinstance(record,dict)
                and record.get('content_digest')==content_hash(card)
                and record.get('context_digest')==image_context(card)
                and record.get('validation_digest')==image_validation_digest(card)
                and record.get('image_digest')==image_final_digest(card)
                and record.get('artifact_digest')==image_approval_artifact_digest(card))

def image_plan_confirmed(card):
    return bool(not image_plan_errors(card) and (card.get('image_plan_approval') or {}).get('context')==image_plan_context(card))

def image_prompt(card):
    parts=[f"그림 유형: {card.get('image_mode','').upper()}",'그림에 문자·숫자·물리량·수식·단위·라벨을 절대 그리지 않는다. 글씨는 별도 편집 레이어에서 처리한다.']
    if card.get('image_text_none'): parts.append('[글씨 없음: 별도 글씨 객체를 만들지 않는다.]')
    if card.get('image_mode') in ('strict','mixed'): parts+=['[SVG 규칙]',card.get('image_locked','')]
    if card.get('image_mode') in ('flexible','mixed'): parts+=['[창작 허용 영역]',card.get('image_creative','')]
    return '\n'.join(parts)

def check_subject(subject, cards):
    available={g['id'] for g in subjects(standards())}
    if subject not in available: raise ValueError('시험 과목을 선택해 주세요.')
    for c in cards:
        code=c.get('standard','')
        if code and (code not in standards() or subject_id(code)!=subject):
            raise ValueError(f"{c['no']}번: {code}는 선택한 과목의 성취기준이 아닙니다.")

def cards_digest(cards):
    rows=[]
    for c in cards:
        row=(c['no'],content_hash(c),bool(c['review_approved']))
        if calculation_override_active(c):
            row+=(c['calculation_override'].get('validation_digest'),)
        rows.append(row)
    return hashlib.sha256(json.dumps(rows).encode()).hexdigest()

def export_digest(exam):
    """Hash exactly the exam text, order and finalized artwork that affect delivered files."""
    exported=[]
    for card in exam['cards']:
        fields={key:card.get(key,'') for key in COLS.values() if key!='no'}
        fields['no']=card['no']; fields['data_kind']=card.get('data_kind','none')
        if calculation_override_active(card):
            fields['calculation_override']={k:card['calculation_override'].get(k) for k in ('content_digest','validation_digest','approved_at')}
        image=card.get('image') or {}
        if card.get('data_kind')=='image' and image:
            fields['final_image']={'final_sha256':image.get('final_sha256'),'current_digest':image_final_digest(card),
                                   'final_digest':image.get('final_digest'),'source_sha256':image.get('sha256')}
            if image_override_active(card):
                fields['final_image']['approval_override']={k:image['approval_override'].get(k) for k in ('content_digest','context_digest','validation_digest','image_digest','artifact_digest','approved_at')}
        else:
            fields['final_image']=None
        if card.get('materials') is not None and not material_array_matches_legacy(card,card['materials']):
            exported_materials=[]
            for block in card['materials']:
                item={k:block.get(k,'') for k in ('id','kind','label','content')}
                if block.get('kind')=='image':
                    projected,_=material_projection(card,block['id'],require_image=True)
                    image=block.get('image') or {}
                    item['final_image']={'final_sha256':image.get('final_sha256'),'current_digest':image_final_digest(projected),
                                         'final_digest':image.get('final_digest'),'source_sha256':image.get('sha256')}
                    if image_override_active(projected):
                        item['final_image']['approval_override']={k:image['approval_override'].get(k) for k in ('content_digest','context_digest','validation_digest','image_digest','artifact_digest','approved_at')}
                exported_materials.append(item)
            fields['materials']=exported_materials
        exported.append(fields)
    payload={'title':exam.get('title',''),'subject':exam.get('subject',''),'cards':exported}
    return hashlib.sha256(json.dumps(payload,sort_keys=True,ensure_ascii=False,separators=(',',':')).encode()).hexdigest()

def cards_approved(exam):
    approval=exam.get('card_approval')
    return bool(approval and approval.get('digest')==cards_digest(exam['cards']))

def allocate_points(cards):
    """Largest-remainder allocation in tenths. Difficulty bands stay ordered."""
    from fractions import Fraction
    if not cards or any(c['difficulty'] not in ('하','중','상') for c in cards):
        raise ValueError('모든 문항의 상·중·하 난이도를 먼저 입력해 주세요.')
    weights={}; base={'하':100,'중':125,'상':150}
    for level in base:
        group=[i for i,c in enumerate(cards) if c['difficulty']==level]
        for rank,i in enumerate(group):
            # Small monotonic increments within a band; do not change question order.
            weights[i]=Fraction(base[level])+ (Fraction(10*rank,len(group)-1)-5 if len(group)>1 else 0)
    total=sum(weights.values())
    exact={i:1000*w/total for i,w in weights.items()}
    units={i:int(v) for i,v in exact.items()}
    for i in sorted(exact,key=lambda i:(exact[i]-units[i],weights[i],-i),reverse=True)[:1000-sum(units.values())]: units[i]+=1
    for i,c in enumerate(cards):
        points=f'{units[i]/10:g}'
        if c['points']!=points: c['review_approved']=False
        c['points']=points
    return cards

def normalize(card, no):
    c = blank_card(no)
    for k in c:
        if k in card:
            c[k] = card[k]
    c['no'] = no
    if not re.fullmatch('[a-f0-9]{16}', str(c['id'])):
        raise ValueError('문항 ID 오류')
    for k in TEXT_FIELDS:
        c[k] = str(c[k] or '').strip()
        if len(c[k]) > 50000:
            raise ValueError('문항 내용이 너무 깁니다.')
    c['authoring_mode']=normalize_authoring_mode(c.get('authoring_mode','legacy'))
    c['materials_request']=normalize_materials_request(c.get('materials_request',''))
    c['material_request_kinds']=normalize_material_request_kinds(c.get('material_request_kinds',[]))
    c['material_request_image_mode']=normalize_material_request_image_mode(c.get('material_request_image_mode',''))
    validate_material_request_preferences(c['material_request_kinds'],c['material_request_image_mode'])
    if c['type'] and c['type'] not in TYPES:
        raise ValueError(f"{no}번: 문항 유형은 네 가지 지정 유형만 가능합니다.")
    if c['difficulty'] not in ('','하','중','상'):
        raise ValueError('난이도는 하·중·상 중 하나여야 합니다.')
    if c['data_kind'] not in ('none','image','text','table'):
        raise ValueError('자료 유형 오류')
    if c['image_mode'] not in ('',*IMAGE_MODES): raise ValueError('그림 유형은 strict / mixed / flexible 중 하나여야 합니다.')
    if type(c['ai_requested']) is not bool: raise ValueError('AI 생성 요청은 Yes 또는 No여야 합니다.')
    if type(c['image_text_none']) is not bool: raise ValueError('글씨 없음 설정은 Yes 또는 No여야 합니다.')
    from .labels import normalize_labels
    c['image_labels']=normalize_labels(c['image_labels'])
    from .vectors import normalize_scene
    c['image_scene']=normalize_scene(c['image_scene'])
    c['labels_initialized']=bool(c['labels_initialized'])
    if c['image_plan_version'] not in (1,2): raise ValueError('그림 계획 버전 오류')
    normalize_materials(c)
    a = str(c['answer'] or '').strip()
    c['answer'] = '①②③④⑤'.index(a) + 1 if a in '①②③④⑤' and a else (int(a) if a.isdigit() else '')
    if c['answer'] not in ('',1,2,3,4,5):
        raise ValueError('정답은 1~5여야 합니다.')
    if c['points']:
        try:
            p = float(c['points'])
        except ValueError:
            raise ValueError('배점은 숫자여야 합니다.')
        if not math.isfinite(p) or p <= 0 or p > 100:
            raise ValueError('배점은 0 초과 100 이하여야 합니다.')
    if not isinstance(c['checks'], list) or len(c['checks']) > 30:
        raise ValueError('검산 목록 오류')
    if any(not isinstance(item,dict) for item in c['checks']):
        raise ValueError('검산식은 이름·왼쪽 식·오른쪽 식을 가진 항목이어야 합니다.')
    # Client-provided approvals are never authoritative; the server restores a
    # still-current record from the previously saved card after normalization.
    c['calculation_override']=None
    return c

@lru_cache(maxsize=1)
def standards():
    w = openpyxl.load_workbook(TEMPLATE, data_only=True, read_only=True)
    out = {}
    for row in w.worksheets[1].iter_rows(min_row=3, values_only=True):
        if len(row)>4 and row[2]:
            out[str(row[2])] = {'name':str(row[3] or ''),'note':str(row[4] or ''),'subject':subject_id(str(row[2]))}
    w.close()
    return out

def import_cards(raw):
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        if sum(x.file_size for x in z.infolist()) > 60_000_000:
            raise ValueError('압축 해제 크기가 너무 큰 엑셀입니다.')
    w = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    if '문항카드 생성' not in w.sheetnames:
        raise ValueError("'문항카드 생성' 시트가 필요합니다.")
    s = w['문항카드 생성']; out = []; catalog = standards()
    headers = {str(s.cell(10,c).value):c for c in range(1,s.max_column+1)}
    for row in s.iter_rows(min_row=11, max_row=110, max_col=max(20,s.max_column),values_only=True):
        if not any(row[c-1] for c in (3,9,10,11,12,13,14)):
            continue
        c = blank_card(len(out)+1)
        for col,k in COLS.items():
            v = row[col-1]
            c[k] = '' if v is None or (isinstance(v,str) and v.startswith('#')) else v
        for label,k in [('내용 영역','concepts'),('해설','explanation')]:
            if label in headers:
                c[k] = row[headers[label]-1] or ''
        for label,k in [('그림 유형','image_mode'),('SVG 고정 영역','image_locked'),('창작 허용 영역','image_creative'),('AI 수정 요청','ai_feedback'),('글씨 프롬프트-물리량','image_text_physics'),('글씨 프롬프트-수식','image_text_math'),('글씨 프롬프트-한글','image_text_korean'),('작성 방식','authoring_mode'),('필요한 자료','materials_request'),('요청 그림 유형','material_request_image_mode')]:
            if label in headers:
                value=row[headers[label]-1]
                if k=='authoring_mode' and not value: continue
                if k=='material_request_image_mode':
                    c[k]=value if value is not None and not (isinstance(value,str) and not value.strip()) else ''
                else:
                    c[k]=value or ''
        if '필요한 자료 종류 JSON' in headers:
            raw_kinds=row[headers['필요한 자료 종류 JSON']-1]
            if raw_kinds is not None and str(raw_kinds).strip():
                try:
                    c['material_request_kinds']=json.loads(str(raw_kinds))
                except (TypeError,ValueError) as exc:
                    raise ValueError('필요한 자료 종류 JSON을 확인해 주세요. image, text, table 배열이어야 합니다.') from exc
        if '글씨 객체 JSON' in headers:
            c['image_labels']=json.loads(row[headers['글씨 객체 JSON']-1] or '[]')
        if '글씨 없음' in headers: c['image_text_none']=str(row[headers['글씨 없음']-1] or '').strip().lower()=='yes'
        if '글씨 배치 완료' in headers: c['labels_initialized']=str(row[headers['글씨 배치 완료']-1]).lower()=='yes'
        c['image_plan_version']=int(row[headers['그림 계획 버전']-1] or 1) if '그림 계획 버전' in headers else 1
        if '자료 블록 JSON' in headers:
            raw_materials=row[headers['자료 블록 JSON']-1]
            if raw_materials:
                materials=json.loads(str(raw_materials))
                if not isinstance(materials,list): raise ValueError('자료 블록 JSON은 배열이어야 합니다.')
                # A workbook does not carry the underlying asset files or an approval authority.
                for block in materials:
                    if isinstance(block,dict) and block.get('kind')=='image':
                        block['image']=None; block['image_plan_approval']=None; block['image_scene']=None
                c['materials']=materials
        c['ai_requested']=not bool(c['question'])
        if 'AI 생성 요청' in headers:
            raw_flag=str(row[headers['AI 생성 요청']-1] or '').strip().lower()
            if raw_flag not in ('yes','no',''): raise ValueError('AI 생성 요청 열은 Yes 또는 No로 입력해 주세요.')
            c['ai_requested']=raw_flag=='yes'
        for key in ('stem','material','question','statements','choices'):
            if str(c[key]).strip() in ('없음.','없음','None'):
                c[key]=''
        old_type=str(c['type']).strip()
        canonical=next((t for t in TYPES if t.replace(' ','')==old_type.replace(' ','')),None)
        if canonical:
            c['type']=canonical
        elif old_type=='객관식':
            question=str(c['question'])
            if '짝지' in question or '짝을' in question: c['type']='짝짓기형'
            elif c['statements'] and ('있는 대로' in question or '옳은 것만' in question): c['type']='합답형'
            elif '들어갈' in question or '빈칸' in question: c['type']='불완전 문장형'
            else: c['type']='단일 정답형'
            c['review_note']=f"기존 유형 '{old_type}'을 문항 구조에 따라 '{c['type']}'으로 임시 대응했습니다. 교사 확인이 필요합니다."
        if c['standard'] in catalog:
            c['standard_name'] = catalog[c['standard']]['name']
            c['standard_note'] = catalog[c['standard']]['note']
        m = str(c['material']).strip()
        c['data_kind'] = 'none' if m in ('','없음.','없음') else ('text' if m.startswith('(글상자)') else ('table' if '|' in m else 'image'))
        if '자료 종류' in headers: c['data_kind']=str(row[headers['자료 종류']-1] or 'none')
        c['calculation'] = '$' in str(c['choices']) or '$' in str(c['question'])
        if str(c['no']).strip() not in (str(len(out)+1),str(float(len(out)+1))):
            raise ValueError(f"문항 번호는 1번부터 빈 번호 없이 이어져야 합니다. {len(out)+1}번 위치를 확인해 주세요.")
        out.append(normalize(c,len(out)+1))
    w.close()
    if not out:
        raise ValueError('입력된 문항이 없습니다. 성취기준 또는 핵심 아이디어를 작성해 주세요.')
    return out

def export_xlsx(cards, path):
    """Edit only card cells in the supplied OOXML package; preserve validations/styles/other sheets."""
    ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    extra = {21:('ai_requested','AI 생성 요청'),22:('image_mode','그림 유형'),23:('image_locked','SVG 고정 영역'),24:('image_creative','창작 허용 영역'),25:('ai_feedback','AI 수정 요청'),26:('image_text_physics','글씨 프롬프트-물리량'),27:('image_text_math','글씨 프롬프트-수식'),28:('image_text_korean','글씨 프롬프트-한글'),29:('image_labels','글씨 객체 JSON')}
    extra.update({30:('labels_initialized','글씨 배치 완료'),31:('image_plan_version','그림 계획 버전'),32:('data_kind','자료 종류'),33:('image_text_none','글씨 없음'),34:('materials','자료 블록 JSON'),
                  35:('authoring_mode','작성 방식'),36:('materials_request','필요한 자료'),
                  37:('material_request_kinds','필요한 자료 종류 JSON'),38:('material_request_image_mode','요청 그림 유형')})
    columns = {**COLS,**{col:pair[0] for col,pair in extra.items()}}
    ET.register_namespace('', ns)
    with zipfile.ZipFile(TEMPLATE) as z:
        root = ET.fromstring(z.read('xl/worksheets/sheet1.xml'))
        rows = {int(r.get('r')):r for r in root.find(f'{{{ns}}}sheetData')}
        for r in range(11,111):
            row = rows[r]
            styles={cell.get('r'):cell.get('s','0') for cell in row}
            for cell in list(row):
                col = re.sub(r'\d','',cell.get('r',''))
                if col in [openpyxl.utils.get_column_letter(c) for c in columns]:
                    row.remove(cell)
            if r-11 >= len(cards):
                continue
            c = cards[r-11]
            for col,k in columns.items():
                value = c.get(k,'')
                if k in ('ai_requested','labels_initialized','image_text_none'): value='Yes' if value else 'No'
                if k=='image_labels': value=json.dumps(value or [],ensure_ascii=False)
                if k=='material_request_kinds': value=json.dumps(value or [],ensure_ascii=False)
                if k=='materials':
                    blocks=c.get('materials')
                    if blocks is None: value=''
                    else:
                        serialized=[]
                        for block in blocks:
                            item=copy.deepcopy(block)
                            if item.get('kind')=='image':
                                item['image']=None; item['image_plan_approval']=None; item['image_scene']=None
                            serialized.append(item)
                        value=json.dumps(serialized,ensure_ascii=False)
                address=f'{openpyxl.utils.get_column_letter(col)}{r}'
                cell = ET.SubElement(row, f'{{{ns}}}c', {'r':address,'t':'inlineStr','s':styles.get(address,styles.get(f'J{r}','0'))})
                if k in ('no','answer','estimate','discrimination') and isinstance(value,(int,float)):
                    cell.set('t','n'); ET.SubElement(cell,f'{{{ns}}}v').text=str(value)
                else:
                    ET.SubElement(ET.SubElement(cell,f'{{{ns}}}is'),f'{{{ns}}}t').text = str(value)
            row[:] = sorted(row, key=lambda c:openpyxl.utils.column_index_from_string(re.sub(r'\d','',c.get('r'))))
        for col,label in [('S','내용 영역'),('T','해설')]+[(openpyxl.utils.get_column_letter(col),pair[1]) for col,pair in extra.items()]:
            row=rows[10]
            cell=next((c for c in row if c.get('r')==col+'10'),None)
            if cell is None: cell=ET.SubElement(row,f'{{{ns}}}c',r=col+'10')
            for child in list(cell): cell.remove(child)
            cell.set('t','inlineStr')
            ET.SubElement(ET.SubElement(cell,f'{{{ns}}}is'),f'{{{ns}}}t').text=label
        dimension=root.find(f'{{{ns}}}dimension')
        if dimension is not None: dimension.set('ref','A1:AL110')
        # Values are materialized, avoiding stale cached lookup formulas.
        with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as dst:
            for item in z.infolist():
                dst.writestr(item, ET.tostring(root,encoding='utf-8',xml_declaration=True) if item.filename=='xl/worksheets/sheet1.xml' else z.read(item.filename))

def safe_expr(source):
    import sympy as sp
    if len(source)>300:
        raise ValueError('검산식이 너무 깁니다.')
    tree = ast.parse(source, mode='eval')
    if sum(1 for _ in ast.walk(tree))>100:
        raise ValueError('검산식이 너무 복잡합니다.')
    def walk(n):
        if isinstance(n,ast.Constant) and type(n.value) in (int,float) and abs(n.value)<1e12:
            return sp.Rational(str(n.value))
        if isinstance(n,ast.Name) and re.fullmatch('[A-Za-z][A-Za-z0-9]{0,15}',n.id):
            return {'pi':sp.pi}.get(n.id,sp.Symbol(n.id,positive=True))
        if isinstance(n,ast.UnaryOp) and isinstance(n.op,(ast.USub,ast.UAdd)):
            return -walk(n.operand) if isinstance(n.op,ast.USub) else walk(n.operand)
        if isinstance(n,ast.BinOp) and isinstance(n.op,(ast.Add,ast.Sub,ast.Mult,ast.Div,ast.Pow)):
            a,b=walk(n.left),walk(n.right)
            if isinstance(n.op,ast.Pow):
                if not b.is_number or abs(float(b))>10:
                    raise ValueError('지수는 -10~10의 상수만 허용합니다.')
                if a.is_number and (len(str(a))>200 or a.has(sp.zoo,sp.oo,sp.nan)):
                    raise ValueError('검산 숫자의 크기가 너무 큽니다.')
                return a**b
            return {ast.Add:lambda:a+b,ast.Sub:lambda:a-b,ast.Mult:lambda:a*b,ast.Div:lambda:a/b}[type(n.op)]()
        if isinstance(n,ast.Call) and isinstance(n.func,ast.Name) and n.func.id in ('sqrt','sin','cos','tan','Abs') and len(n.args)==1 and not n.keywords:
            return getattr(sp,n.func.id)(walk(n.args[0]))
        raise ValueError('허용되지 않는 검산식입니다.')
    return walk(tree.body)

def verify_checks(checks):
    import sympy as sp
    out=[]
    for check in checks:
        try:
            if not isinstance(check,dict): raise ValueError('검산 항목 형식 오류')
            diff=sp.simplify(safe_expr(str(check['lhs']))-safe_expr(str(check['rhs'])))
            ok=diff == 0
            out.append({'label':str(check.get('label','검산')),'ok':bool(ok),'detail':str(diff)})
        except Exception as e:
            label=check.get('label','검산') if isinstance(check,dict) else '검산'
            out.append({'label':str(label),'ok':False,'detail':str(e)})
    return out

def validate(exam, phase='export'):
    errors=[]; issues=[]; warnings=[]; cards=exam['cards']; total=0
    def issue(code, message, card=None, field=None, stage='review', material_id=None):
        errors.append(message)
        issues.append({'code':code,'card_id':card.get('id') if card else None,'material_id':material_id,'field':field,
                       'stage':stage,'severity':'error','message':message})
    if not cards: issue('cards_required','문항을 추가해 주세요.',stage='design')
    for c in cards:
        n=c['no']
        for k,label in [('standard','성취기준'),('type','문항 유형'),('difficulty','난이도'),('idea','핵심 아이디어'),('question','발문'),('choices','선택지'),('answer','정답'),('points','배점'),('explanation','해설')]:
            if not c.get(k): issue(f'{k}_required',f'{n}번: {label} 누락',c,k,'review')
        nums = re.findall('[①②③④⑤]', c['choices'])
        if nums != list('①②③④⑤'): issue('choices_format','%s번: 선택지는 ①~⑤ 순서로 한 번씩 필요합니다.'%n,c,'choices','review')
        if c['type']=='합답형' and not c['statements']: issue('statements_required',f'{n}번: 합답형 보기가 없습니다.',c,'statements','review')
        if not c['review_approved']: issue('review_required',f'{n}번: 문항 검토 확인이 필요합니다.',c,'review_approved','review')
        if c['calculation']:
            input_errors=calculation_input_errors(c)
            if input_errors:
                issue('calculation_input_required',f'{n}번: '+ ' / '.join(input_errors),c,'checks','review')
            else:
                results=verify_checks(c['checks'])
                if any(not x['ok'] for x in results):
                    if calculation_override_active(c):
                        warnings.append(f'{n}번: 계산 불일치를 교사 확인 후 승인했습니다.')
                    else:
                        issue('calculation_check_required',f'{n}번: 계산식이 서로 일치하지 않습니다. 확인 후 다시 검산하거나 “그래도 승인”을 선택해 주세요.',c,'checks','review')
        if phase=='export':
            for block in material_blocks(c):
                if block.get('kind')!='image': continue
                projected,mid=material_projection(c,block['id'],require_image=True)
                plan_errors=image_plan_errors(projected)
                if not plan_errors and not image_plan_confirmed(projected):
                    issue('image_plan_unconfirmed',f"{n}번 {block.get('label') or '그림'}: 제작 계획을 명시적으로 확인해 주세요.",c,'image_plan_approval','image-confirm',mid)
                im=projected.get('image')
                failed_checks=[x for x in image_check_results((im or {}).get('checks') or []) if not x['ok']]
                image_override=image_override_active(projected)
                if failed_checks and image_override:
                    warnings.append(f"{n}번 {block.get('label') or '그림'}: 그림 제약 불일치를 교사 확인 후 승인했습니다.")
                if failed_checks and not image_override:
                    issue('image_check_failed',f"{n}번 {block.get('label') or '그림'}: 그림 제약 검사 결과를 확인하고 최종 확정해 주세요.",c,'image','image',mid)
                if (not im or not im.get('approved') or im.get('context')!=image_context(projected)
                        or im.get('final_digest')!=image_final_digest(projected) or not im.get('final_file')
                        or (failed_checks and not image_override)):
                    issue('image_approval_required',f"{n}번 {block.get('label') or '그림'}: 그림과 글씨를 최종 확정해 주세요.",c,'image','image',mid)
        total+=float(c['points'] or 0)
    if cards and sum(c['type']=='합답형' for c in cards)>len(cards)/2:
        issue('combination_ratio','합답형 비율이 전체의 50%를 초과합니다.',stage='design')
    try: check_subject(exam.get('subject') or infer_subject(cards),cards)
    except ValueError as e: issue('standard_subject_mismatch',str(e),stage='design',field='standard')
    for c in cards:
        for block in material_blocks(c):
            if block.get('kind')!='image': continue
            projected,mid=material_projection(c,block['id'],require_image=True)
            for message in image_plan_errors(projected):
                field='image_mode' if ('유형' in message or '프롬프트' in message) else 'image_locked' if 'SVG' in message or '고정' in message or '같은 요소' in message or '요소ID' in message else 'image_creative'
                issue('image_plan_invalid',f"{c['no']}번 {block.get('label') or '그림'}: {message}",c,field,'image-plan',mid)
    if abs(total-100)>0.001: issue('score_total',f'현재 배점 합은 {total:g}점입니다. 배점 패널에서 합계 100점을 맞춰 주세요.',stage='score',field='points')
    if phase=='export' and not cards_approved(exam): issue('approval_required','전체 출제안 승인이 필요합니다. 승인 후 그림 제작과 출력을 진행할 수 있습니다.',stage='approval')
    return {'errors':errors,'issues':issues,'warnings':warnings,'points':round(total,3),'ready':not errors}
